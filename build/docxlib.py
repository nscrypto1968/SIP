"""Formatting engine for the Rosy Blue Securities Black Book.

Every value in the master prompt's PAGE SETUP / TEXT FORMATS / DATA TABLE FORMAT
specification is applied here as direct formatting, to the twip.
"""
import re
from docx import Document
from docx.shared import Pt, Twips, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement

FONT = "Times New Roman"
NAVY = "1F3A73"
MID = "4F81BD"
LIGHT = "A9BCDC"
HDR_FILL = "D9E2F3"
TOT_FILL = "F2F2F2"


# ---------------------------------------------------------------- low level
def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e


def set_spacing(p, before=None, after=None, line=None, rule="auto"):
    pf = p.paragraph_format
    if before is not None:
        pf.space_before = Twips(before)
    if after is not None:
        pf.space_after = Twips(after)
    if line is not None:
        pPr = p._p.get_or_add_pPr()
        sp = pPr.find(qn("w:spacing"))
        if sp is None:
            sp = el("w:spacing")
            pPr.append(sp)
        sp.set(qn("w:line"), str(line))
        sp.set(qn("w:lineRule"), rule)


def keep_next(p, on=True):
    pPr = p._p.get_or_add_pPr()
    k = pPr.find(qn("w:keepNext"))
    if on and k is None:
        pPr.append(el("w:keepNext"))


def outline(p, lvl):
    pPr = p._p.get_or_add_pPr()
    pPr.append(el("w:outlineLvl", val=lvl))


def page_break_before(p):
    pPr = p._p.get_or_add_pPr()
    pPr.append(el("w:pageBreakBefore"))


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(el("w:shd", val="clear", color="auto", fill=fill))


def cell_margins(table, top=40, bottom=40, left=80, right=80):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        n = el("w:" + side, w=v, type="dxa")
        m.append(n)
    tblPr.append(m)


def table_borders(table, sz=4):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b.append(el("w:" + side, val="single", sz=sz, space=0, color="000000"))
    tblPr.append(b)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(el("w:tblHeader"))


def cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(el("w:cantSplit"))


def vcenter(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(el("w:vAlign", val="center"))


# ---------------------------------------------------------------- runs
PLACEHOLDER = re.compile(r"(\[ADD:[^\]]*\])")
TOKEN = re.compile(r"(\*\*.+?\*\*|__.+?__|\[ADD:[^\]]*\])", re.S)


def add_runs(p, text, size=12, bold=False, italic=False, underline=False):
    """Mini-markup: **bold**, __italic__, [ADD: ...] -> yellow-highlighted."""
    for part in TOKEN.split(text):
        if not part:
            continue
        b, i, hl = bold, italic, False
        t = part
        if part.startswith("**") and part.endswith("**"):
            t, b = part[2:-2], True
        elif part.startswith("__") and part.endswith("__"):
            t, i = part[2:-2], True
        elif PLACEHOLDER.fullmatch(part):
            hl = True
        r = p.add_run(t)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.bold = b
        r.italic = i
        r.underline = underline
        rPr = r._r.get_or_add_rPr()
        rf = rPr.find(qn("w:rFonts"))
        if rf is None:
            rf = el("w:rFonts")
            rPr.insert(0, rf)
        for a in ("ascii", "hAnsi", "cs", "eastAsia"):
            rf.set(qn("w:" + a), FONT)
        if hl:
            rPr.append(el("w:highlight", val="yellow"))
    return p


# ---------------------------------------------------------------- builder
class Report:
    def __init__(self):
        self.doc = Document()
        self._setup()
        self.pending_break = False

    # ---- document setup -------------------------------------------------
    def _setup(self):
        d = self.doc
        st = d.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(12)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = el("w:rFonts")
            rpr.insert(0, rf)
        for a in ("ascii", "hAnsi", "cs", "eastAsia"):
            rf.set(qn("w:" + a), FONT)
        rpr.append(el("w:lang", val="en-IN", eastAsia="en-IN", bidi="ar-SA"))
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.space_before = Pt(0)

        sec = d.sections[0]
        sec.page_width = Twips(11906)
        sec.page_height = Twips(16838)
        sec.top_margin = Twips(1300)
        sec.bottom_margin = Twips(1300)
        sec.left_margin = Twips(1440)
        sec.right_margin = Twips(1440)
        sec.header_distance = Twips(708)
        sec.footer_distance = Twips(500)
        sec.different_first_page_header_footer = True

        sectPr = sec._sectPr
        # page border: 2.25pt black, offset from page, space 24, all sides
        pgb = OxmlElement("w:pgBorders")
        pgb.set(qn("w:offsetFrom"), "page")
        for side in ("top", "left", "bottom", "right"):
            pgb.append(el("w:" + side, val="single", sz=18, space=24, color="000000"))
        sectPr.append(pgb)

        # footer: centred PAGE field, TNR 11pt
        f = sec.footer
        p = f.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        r.font.name = FONT
        r.font.size = Pt(11)
        fld = el("w:fldChar", fldCharType="begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        end = el("w:fldChar", fldCharType="end")
        r._r.append(fld)
        r._r.append(instr)
        r._r.append(end)
        # first-page footer stays empty (cover carries no number)

    # ---- generic paragraph ---------------------------------------------
    def _p(self):
        p = self.doc.add_paragraph()
        if self.pending_break:
            page_break_before(p)
            self.pending_break = False
        return p

    def newpage(self):
        self.pending_break = True

    # ---- spec'd paragraph types ----------------------------------------
    def body(self, text, size=12):
        p = self._p()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_spacing(p, after=100, line=360)
        add_runs(p, text, size=size)
        return p

    def interpretation(self, text):
        return self.body("**Interpretation.** " + text)

    def chapter(self, label, title):
        self.newpage()
        p = self._p()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(p, after=80)
        add_runs(p, label, size=14, bold=True, underline=True)
        outline(p, 0)
        q = self.doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(q, after=260)
        add_runs(q, title, size=14, bold=True, underline=True)
        outline(q, 0)

    def heading_centre(self, text, after=300):
        p = self._p()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(p, after=after)
        add_runs(p, text, size=14, bold=True, underline=True)
        outline(p, 0)
        return p

    def section(self, text):
        p = self._p()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_spacing(p, before=200, after=110)
        keep_next(p)
        add_runs(p, text.upper(), size=12, bold=True)
        outline(p, 1)
        return p

    def sub(self, text):
        p = self._p()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_spacing(p, before=120, after=70)
        keep_next(p)
        add_runs(p, text, size=12, bold=True, italic=True)
        outline(p, 2)
        return p

    def bullet(self, lead, text=None):
        p = self._p()
        try:
            p.style = self.doc.styles["List Paragraph"]
        except KeyError:
            pass
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.left_indent = Twips(540)
        pf.first_line_indent = Twips(-300)
        set_spacing(p, after=50, line=340)
        r = p.add_run("\u2022\t")
        r.font.name = FONT
        r.font.size = Pt(12)
        if text is None:
            add_runs(p, lead, size=12)
        else:
            add_runs(p, "**" + lead + ":** " + text, size=12)
        return p

    def caption(self, text):
        p = self._p()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(p, before=140, after=70)
        keep_next(p)
        add_runs(p, text, size=11, bold=True)
        return p

    def source(self, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(p, before=40, after=140)
        add_runs(p, text, size=10, italic=True)
        return p

    def formula(self, text):
        p = self._p()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(p, before=60, after=100, line=300)
        add_runs(p, text, size=12, italic=True)
        return p

    def image(self, path, width_in=5.7):
        p = self._p()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(p, before=0, after=0)
        keep_next(p)
        p.add_run().add_picture(path, width=Inches(width_in))
        return p

    def spacer(self, pts=6):
        p = self._p()
        set_spacing(p, after=int(pts * 20))
        return p

    # ---- tables ----------------------------------------------------------
    def table(self, headers, rows, widths=None, size=None, total_rows=(),
              align=None, total_width=9000):
        ncols = len(headers)
        if widths is None:
            base = total_width // ncols
            widths = [base] * ncols
            widths[-1] = total_width - base * (ncols - 1)
        assert sum(widths) == total_width, (sum(widths), total_width)
        if size is None:
            size = 9.5 if (ncols >= 6) else 10.5
        if align is None:
            align = ["left"] + ["center"] * (ncols - 1)

        t = self.doc.add_table(rows=0, cols=ncols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        tblPr.append(el("w:tblW", w=total_width, type="dxa"))
        tblPr.append(el("w:tblLayout", type="fixed"))
        table_borders(t)
        cell_margins(t)
        # grid
        grid = t._tbl.find(qn("w:tblGrid"))
        if grid is not None:
            t._tbl.remove(grid)
        grid = OxmlElement("w:tblGrid")
        for w in widths:
            grid.append(el("w:gridCol", w=w))
        t._tbl.insert(1, grid)

        def fill_row(cells, values, bold=False, fillc=None):
            for i, (c, v) in enumerate(zip(cells, values)):
                c.width = Twips(widths[i])
                tcPr = c._tc.get_or_add_tcPr()
                tcPr.append(el("w:tcW", w=widths[i], type="dxa"))
                vcenter(c)
                if fillc:
                    shade(c, fillc)
                p = c.paragraphs[0]
                a = align[i] if i < len(align) else "center"
                p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT,
                               "center": WD_ALIGN_PARAGRAPH.CENTER,
                               "right": WD_ALIGN_PARAGRAPH.RIGHT}[a]
                set_spacing(p, before=0, after=0, line=250)
                keep_next(p)
                add_runs(p, str(v), size=size, bold=bold)

        hr = t.add_row()
        repeat_header(hr)
        cant_split(hr)
        fill_row(hr.cells, headers, bold=True, fillc=HDR_FILL)
        for idx, row in enumerate(rows):
            r = t.add_row()
            cant_split(r)
            istot = idx in total_rows
            fill_row(r.cells, row, bold=istot, fillc=TOT_FILL if istot else None)
        return t

    def save(self, path):
        self.doc.save(path)
