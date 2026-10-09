#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Builds the Rosy Blue Securities MMS Black Book .docx."""
import os, sys
sys.path.insert(0, os.path.dirname(__file__))

import compute as C
from compute import inr, pct
import charts as G
from docxlib import Report, add_runs, set_spacing, keep_next, page_break_before, el, shade
from docx.shared import Pt, Twips, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(os.path.dirname(HERE), "Rosy_Blue_Securities_Black_Book.docx")
TODAY = "9 October 2026"
AS_ON = "25 September 2026"

R = Report()
d = R.doc

SRC_NSE = "Source: NSE India, circulars and market-statistics pages; checked 9 October 2026."
SRC_SEBI = "Source: SEBI circulars and master circulars; checked 9 October 2026."
SRC_AUTH = "Source: Author's computation; illustrative figures."
SRC_PL_FY26 = ("Source: Pine Labs Limited, audited consolidated financial results for the year "
               "ended 31 March 2026 (board meeting of 25 May 2026); checked 9 October 2026.")
SRC_PL_Q1 = ("Source: Pine Labs Limited, unaudited consolidated financial results for the quarter "
             "ended 30 June 2026 (board meeting of 28 July 2026); checked 9 October 2026.")
SRC_PL_RHP = ("Source: Pine Labs Limited, Red Herring Prospectus, October 2025, restated "
              "consolidated financial information; checked 9 October 2026.")
SRC_COMP = "Source: Compiled by the author."

# ============================================================ FRONT MATTER
# ---- 1. Cover page
def cover():
    def line(text, size, after, before=None):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(p, before=before, after=after, line=276)
        add_runs(p, text, size=size, bold=True)
        return p
    line("BLACK BOOK PROJECT REPORT ON", 14, 160)
    line("A Study of Equity Dealing, the Futures and Options Market and "
         "Trading Terminals at Rosy Blue Securities Pvt. Ltd.:", 14, 40)
    line("WITH A COMPANY ANALYSIS OF PINE LABS LIMITED", 14, 40)
    line("Submitted in partial fulfilment for the award of the degree of", 12, 200)
    line("MASTER OF MANAGEMENT STUDIES", 13, 60)
    line("FINANCE", 13, 260)
    line("UNIVERSITY OF MUMBAI", 13, 260)
    line("Submitted by", 12, 100)
    line("[ADD: STUDENT NAME]", 13, 60)
    line("ROLL NO: [ADD: ROLL NO]", 14, 260)
    line("[ADD: ACADEMIC YEARS, e.g. 2025-2027]", 14, 260)
    line("Under The Guidance of", 14, 100)
    line("[ADD: NAME OF PROJECT GUIDE]", 13, 700)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, after=120, line=276)
    add_runs(p, "[ADD: COLLEGE LOGO IMAGE HERE \u2014 2.08 in wide \u00d7 1.49 in high]",
             size=11, italic=True)
    line("[ADD: NAME OF INSTITUTE]", 14, 40)
    line("[ADD: CITY AND PIN CODE]", 14, 0)


# ---- 2. Certificate
def certificate():
    R.newpage()
    p = R._p()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, after=120)
    add_runs(p, "[ADD: COLLEGE LETTERHEAD IMAGE HERE \u2014 5.3 in wide \u00d7 1.18 in high]",
             size=11, italic=True)
    q = d.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(q, after=260)
    add_runs(q, "CERTIFICATE", size=14, bold=True, underline=True)

    b = d.add_paragraph()
    b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_spacing(b, after=240, line=400)
    add_runs(b,
        "This is to certify that the summer internship project report titled "
        "**\u201cA Study of Equity Dealing, the Futures and Options Market and Trading Terminals "
        "at Rosy Blue Securities Pvt. Ltd.: with a Company Analysis of Pine Labs Limited\u201d** "
        "submitted in partial fulfilment of the requirements for the Master of Management Studies "
        "degree examination of the University of Mumbai by **[ADD: STUDENT NAME]**, "
        "Roll No. **[ADD: ROLL NO]**, is a record of bonafide research work carried out by the "
        "student during the period from [ADD: INTERNSHIP START DATE] to "
        "[ADD: INTERNSHIP END DATE] under my guidance and supervision, and has been found "
        "satisfactory. The work presented in this report is original and has not been submitted "
        "earlier for the award of any degree, diploma or associateship of any other university "
        "or institute.")

    t = d.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    from docxlib import table_borders, cell_margins
    t._tbl.tblPr.append(el("w:tblW", w=2400, type="dxa"))
    t._tbl.tblPr.append(el("w:tblLayout", type="fixed"))
    _g = t._tbl.find(qn("w:tblGrid"))
    if _g is not None:
        t._tbl.remove(_g)
    from docx.oxml import OxmlElement as _OE
    _g = _OE("w:tblGrid"); _g.append(el("w:gridCol", w=2400)); t._tbl.insert(1, _g)
    table_borders(t, sz=4)
    cell_margins(t, 80, 80, 80, 80)
    c = t.cell(0, 0)
    c.width = Twips(2400)
    cp = c.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(cp, before=60, after=60)
    add_runs(cp, "COLLEGE STAMP HERE", size=10)

    sp = d.add_paragraph()
    set_spacing(sp, after=500)

    st = d.add_table(rows=1, cols=2)
    st.alignment = WD_TABLE_ALIGNMENT.CENTER
    st.autofit = False
    st._tbl.tblPr.append(el("w:tblW", w=9000, type="dxa"))
    st._tbl.tblPr.append(el("w:tblLayout", type="fixed"))
    _g2 = st._tbl.find(qn("w:tblGrid"))
    if _g2 is not None:
        st._tbl.remove(_g2)
    from docx.oxml import OxmlElement as _OE2
    _g2 = _OE2("w:tblGrid")
    _g2.append(el("w:gridCol", w=4500)); _g2.append(el("w:gridCol", w=4500))
    st._tbl.insert(1, _g2)
    for i, (name, role) in enumerate([("[ADD: NAME OF DIRECTOR]", "DIRECTOR"),
                                      ("[ADD: NAME OF PROJECT GUIDE]", "PROJECT GUIDE")]):
        cell = st.cell(0, i)
        cell.width = Twips(4500)
        cell._tc.get_or_add_tcPr().append(el("w:tcW", w=4500, type="dxa"))
        p0 = cell.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(p0, after=40)
        add_runs(p0, "-" * 30, size=12)
        for txt in (name, role):
            pp = cell.add_paragraph()
            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_spacing(pp, after=40)
            add_runs(pp, txt, size=12, bold=True)

    e1 = d.add_paragraph()
    e1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(e1, before=500, after=40)
    add_runs(e1, "-" * 30, size=12)
    e2 = d.add_paragraph()
    e2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(e2, "EXTERNAL EXAMINER", size=12, bold=True)


# ---- 3. Experience certificate (placeholder page)
def experience_cert():
    R.newpage()
    R.heading_centre("EXPERIENCE CERTIFICATE")
    p = R._p()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, before=2000, after=0)
    add_runs(p, "[ADD: FULL-PAGE SCANNED IMAGE OF THE INTERNSHIP COMPLETION / EXPERIENCE "
                "CERTIFICATE ISSUED BY ROSY BLUE SECURITIES PVT. LTD. \u2014 if the firm has not "
                "issued one, delete this page and renumber the Contents.]", size=11, italic=True)


# ---- 4. Declaration
def declaration():
    R.newpage()
    R.heading_centre("DECLARATION")
    p = R._p()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_spacing(p, after=600, line=420)
    add_runs(p,
        "I, **[ADD: STUDENT NAME]**, Roll No. **[ADD: ROLL NO]**, a student of the Master of "
        "Management Studies (Finance) programme of the University of Mumbai, hereby declare that "
        "the project report titled **\u201cA Study of Equity Dealing, the Futures and Options "
        "Market and Trading Terminals at Rosy Blue Securities Pvt. Ltd.: with a Company Analysis "
        "of Pine Labs Limited\u201d** has been prepared by me as part of the academic requirements "
        "of the programme. The report is based on my own learning, observation and practical "
        "exposure during the internship undertaken at Rosy Blue Securities Pvt. Ltd. in the "
        "capacity of an Equity Dealer, together with information drawn from publicly available "
        "documents of the exchanges, the Securities and Exchange Board of India and Pine Labs "
        "Limited. Confidential information such as client names, client codes, permanent account "
        "numbers, unique client codes, internal ticket numbers, terminal identifiers, internet "
        "protocol addresses, proprietary configurations and the firm's internal financial data "
        "has not been disclosed; all trading examples given in this report are illustrative and "
        "are identified as such. I further declare that this report has not been submitted "
        "earlier to any other institution or university for the award of any degree, diploma or "
        "other qualification.")

    def sig_line(left, right, after):
        q = d.add_paragraph()
        pf = q.paragraph_format
        pf.tab_stops.add_tab_stop(Twips(9000), WD_TAB_ALIGNMENT.RIGHT)
        set_spacing(q, after=after)
        add_runs(q, left + "\t" + right, size=12)
    sig_line("Name : [ADD: STUDENT NAME]", "Signature", 120)
    sig_line("Roll No. [ADD: ROLL NO]", "[ADD: STUDENT NAME]", 0)


# ---- 5. Acknowledgement
def acknowledgement():
    R.newpage()
    R.heading_centre("ACKNOWLEDGEMENT")
    for text in [
        "I wish to record my sincere gratitude to my project guide, **[ADD: NAME OF PROJECT "
        "GUIDE]**, for the patient supervision, academic direction and constructive criticism "
        "that shaped this report. The guide's insistence that every figure be traced to a "
        "verifiable source, and that every table be followed by an interpretation rather than a "
        "restatement, has improved both the discipline and the usefulness of this work.",

        "I am grateful to my family and friends for their encouragement and forbearance during "
        "the internship and the writing of this report. Their support made it possible to "
        "combine the demands of a live dealing desk, which begins well before the market opens, "
        "with the equally demanding requirements of academic study.",

        "I thank the management of **Rosy Blue Securities Pvt. Ltd.** for granting me the "
        "opportunity to work in the Equity Dealing function, and the dealing-desk seniors, risk "
        "management colleagues and back-office staff who explained processes, corrected my "
        "errors and allowed me to observe live operations. I also thank "
        "**[ADD: NAME OF DIRECTOR]**, Director of the Institute, the faculty members of the "
        "Finance specialisation and the library staff for the academic and infrastructural "
        "support extended throughout the programme.",
    ]:
        p = R._p()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_spacing(p, after=100, line=400)
        add_runs(p, text)
    s = d.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_spacing(s, before=300, after=0)
    add_runs(s, "[ADD: STUDENT NAME]", size=12, bold=True)


# ---- 6. Executive summary
def executive_summary():
    R.newpage()
    R.heading_centre("EXECUTIVE SUMMARY")
    paras = [
        "The Indian secondary market has widened very rapidly over the last five years, and the "
        "operational layer that carries this growth \u2014 the stock broker's dealing desk \u2014 "
        "has received far less academic attention than the valuation and portfolio topics that "
        "dominate finance curricula. As on 31 August 2026 the two depositories together serviced "
        "about " + f"{C.DEMAT_TOTAL:.2f}" + " crore demat accounts, and the National Stock "
        "Exchange reported about " + f"{C.NSE_UNIQUE_INVESTORS:.1f}" + " crore unique registered "
        "investors as of July 2026. Every one of those accounts reaches the market through an "
        "order that must be keyed, risk-checked, routed, matched, cleared and settled within a "
        "defined time cycle. This study examines that process from inside the function, and "
        "pairs it with a fundamental analysis of a recently listed company so that the "
        "operational and analytical sides of a finance role are treated together rather than "
        "separately.",

        "The study was carried out during an internship at Rosy Blue Securities Pvt. Ltd. in the "
        "capacity of an Equity Dealer. The intern observed and, under supervision, participated "
        "in the pre-market routine, order entry and modification on the exchange-provided "
        "terminal and on the member's Computer-to-Computer Link facility, client-wise limit and "
        "margin verification, response to risk-management alerts, square-off discipline and "
        "post-trade reconciliation. The theoretical chapter documents the structure of the "
        "Indian capital market, the trading mechanism, the settlement cycle, the margining "
        "framework in both segments, the regulatory obligations attaching to "
        "Computer-to-Computer Link terminals and the ladder of statutory charges on every "
        "executed trade.",

        "The analytical chapter converts that description into computed evidence. A futures "
        "contract priced by cost of carry on an illustrative underlying at Rs. " +
        inr(C.F_SPOT) + " gives a fair value of Rs. " + inr(C.FUT_FAIR) + ", a basis of " +
        pct(C.FUT_BASIS_PCT, 3) + " per cent over thirty days. A five-day mark-to-market "
        "sequence on one lot produces daily cash flows summing to Rs. " + inr(C.MTM_TOTAL, 0) +
        ", exactly the difference between the entry and the final settlement price, showing "
        "that daily settlement redistributes but does not alter the total result. A thirty-day "
        "call struck at Rs. " + inr(C.O_K, 0) + " prices at Rs. " + inr(C.CALL_PRICE) +
        " under the Black-Scholes model, with a theta of Rs. " + f"{C.BS['theta_c']:.4f}" +
        " per day, and loses roughly two-thirds of its value in the final ten days if the "
        "underlying does not move. A delivery round trip of Rs. 1,00,000 attracts Rs. " +
        inr(C.DEL["total"]) + " of total cost. The company analysis finds that Pine Labs "
        "Limited, listed on 14 November 2025, recorded its first full year of profit in "
        "FY2025-26 with revenue of Rs. " + inr(C.PL_REV[-1]) + " crore, up " +
        pct(C.PL_GROWTH[-1]) + " per cent, and a profit after tax of Rs. " + inr(C.PL_PAT[-1]) +
        " crore against a loss of Rs. " + inr(abs(C.PL_PAT[-2])) + " crore, a swing of Rs. " +
        inr(C.PL_PAT_SWING) + " crore.",

        "The study concludes that competence on a dealing desk rests on three things learned "
        "only by doing: speed and accuracy in order entry, an instinct for a client's margin "
        "position before the risk system raises an alert, and an understanding of the cost "
        "ladder that determines whether a nominally profitable trade is in fact profitable. It "
        "further concludes that the separation of expiry days between the exchanges, the "
        "upfront margin regime and the compressed settlement cycle have collectively made the "
        "role more procedural and less discretionary than a decade ago. Pine Labs demonstrates "
        "genuine operating leverage but trades at a multiple of earnings that leaves little "
        "tolerance for disappointment; the analysis is academic and is not investment advice.",
    ]
    for t in paras:
        p = R._p()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_spacing(p, after=100, line=380)
        add_runs(p, t)


# ---- 7. Contents
CONTENTS = [
    ("Chapter 1", None, "INTRODUCTION", 9),
    (None, "1.1-1.6", "Need, objectives, scope, market overview, challenges, structure", None),
    ("Chapter 2", None, "ORGANISATIONAL BACKGROUND", 13),
    (None, "2.1-2.3", "The firm, its registrations and its services", None),
    (None, "2.4-2.6", "Organisation structure, the dealing desk and the intern's role", None),
    ("Chapter 3", None, "THEORETICAL FRAMEWORK", 18),
    (None, "3.1-3.6", "Financial system, markets, exchanges, participants and regulation", None),
    (None, "3.7-3.9", "Trading mechanism, clearing and settlement, cash-market risk", None),
    (None, "3.10-3.14", "Derivatives, futures, options, margining and strategies", None),
    (None, "3.15-3.17", "NEAT PLUS, CTCL and the back office", None),
    (None, "3.18-3.20", "Transaction costs, investor protection and fintech", None),
    ("Chapter 4", None, "RESEARCH METHODOLOGY", 36),
    (None, "4.1-4.6", "Design, data, period, tools, mapping and ethics", None),
    ("Chapter 5", None, "ANALYSIS AND INTERPRETATION", 39),
    (None, "5.1-5.3", "The dealing day and the market snapshot", None),
    (None, "5.4-5.5", "Order life cycle on NEAT PLUS and CTCL in practice", None),
    (None, "5.6-5.9", "Futures, options, margin and charges worked examples", None),
    (None, "5.10", "Company analysis: Pine Labs Limited", None),
    (None, "5.11-5.13", "Dealer's view, learning outcomes and summary of findings", None),
    ("Chapter 6", None, "CONCLUSION AND SUGGESTIONS", 63),
    (None, "6.1-6.3", "Conclusion, suggestions and learning outcomes", None),
    ("Chapter 7", None, "LIMITATIONS OF THE STUDY", 66),
    ("Chapter 8", None, "BIBLIOGRAPHY", 67),
    (None, "8.1-8.3", "Books, regulations and circulars, websites and filings", None),
    ("Annexure A", None, "GLOSSARY OF TERMS", 69),
    ("Annexure B", None, "FORMULA SHEET", 71),
    ("Annexure C", None, "LIST OF ABBREVIATIONS", 72),
]


def contents():
    R.newpage()
    R.heading_centre("CONTENTS")
    widths = [1250, 1050, 5250, 1450]
    t = d.add_table(rows=0, cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    from docxlib import table_borders, cell_margins, repeat_header, vcenter
    t._tbl.tblPr.append(el("w:tblW", w=9000, type="dxa"))
    t._tbl.tblPr.append(el("w:tblLayout", type="fixed"))
    table_borders(t)
    cell_margins(t)
    grid = t._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        t._tbl.remove(grid)
    from docx.oxml import OxmlElement
    grid = OxmlElement("w:tblGrid")
    for w in widths:
        grid.append(el("w:gridCol", w=w))
    t._tbl.insert(1, grid)

    def row(vals, bold, aligns):
        r = t.add_row()
        for i, (c, v) in enumerate(zip(r.cells, vals)):
            c.width = Twips(widths[i])
            c._tc.get_or_add_tcPr().append(el("w:tcW", w=widths[i], type="dxa"))
            vcenter(c)
            p = c.paragraphs[0]
            p.alignment = aligns[i]
            set_spacing(p, before=0, after=0, line=250)
            add_runs(p, v, size=10, bold=bold)
        return r

    L, Ct = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER
    hr = row(["", "", "CONTENT", "PAGE NO."], True, [L, Ct, Ct, Ct])
    repeat_header(hr)
    for ch, sec, title, page in CONTENTS:
        if ch:
            row([ch, "", title, str(page) if page else ""], True, [L, Ct, L, Ct])
        else:
            row(["", sec, title, ""], False, [L, Ct, L, Ct])
    R.source("Source: Compiled by the author. Page numbers are computed from a layout model "
             "and must be refreshed in Word after the logo, letterhead and "
             "experience-certificate images have been inserted.")


# ============================================================ CHAPTER 1
def chapter1():
    R.chapter("CHAPTER 1", "INTRODUCTION")

    R.section("1.1 Need of the study")
    R.body(
        "Finance education in India devotes considerable attention to valuation, portfolio theory "
        "and corporate finance, and comparatively little to the operational machinery through "
        "which a securities transaction actually occurs. A student may be able to construct a "
        "discounted cash-flow model and yet be unable to explain what happens between the moment "
        "a client telephones an order and the moment the shares appear in that client's demat "
        "account. That gap matters, because most entry-level positions in broking, wealth "
        "management and market operations are concerned with exactly this process. The "
        "**equity dealer** sits where client instruction, regulatory obligation and market "
        "microstructure meet, and the role therefore offers an unusually compressed view of how "
        "the secondary market functions in practice.")
    R.body(
        "The need also arises from the pace of regulatory change. In the last few years the "
        "settlement cycle has been shortened, an upfront margin regime has replaced informal "
        "intraday leverage, the pledge and re-pledge mechanism has altered how client collateral "
        "is held, and the weekly expiry days of the two principal exchanges have been separated "
        "by regulatory direction. Material written even three years ago is therefore unreliable "
        "as a description of current practice. The specific needs addressed are set out below.")
    for lead, txt in [
        ("To connect theory to procedure",
         "to show how price-time priority, cost of carry, option time value and value-at-risk "
         "margining appear in the concrete form of a screen, a limit or an obligation."),
        ("To quantify the cost of trading",
         "to compute the full statutory and transactional cost ladder on illustrative trades, so "
         "that the difference between gross and net outcome is made explicit rather than "
         "assumed."),
        ("To understand trading technology",
         "to examine the exchange-provided terminal and the member's Computer-to-Computer Link "
         "facility, and the pre-trade risk controls and audit obligations attaching to each."),
        ("To apply fundamental analysis to a recent listing",
         "to analyse Pine Labs Limited, a company that listed in November 2025 and reported its "
         "first full year of profit in FY2025-26, as a test of whether standard analytical tools "
         "work on a company with a short listed history."),
    ]:
        R.bullet(lead, txt)

    R.section("1.2 Objectives of the study")
    for lead, txt in [
        ("To study the Indian equity market",
         "its structure, participants, trading mechanism, settlement cycle and the regulatory "
         "framework administered by the Securities and Exchange Board of India and the "
         "exchanges."),
        ("To study the futures and options market",
         "contract design, lot size and expiry conventions, pricing by cost of carry, option "
         "payoff and premium decomposition, and the computation of option sensitivities."),
        ("To examine the NEAT PLUS trading terminal",
         "its market watch, order-entry, order-book and trade-book functions, and the order "
         "modification and cancellation facilities used by a dealer during market hours."),
        ("To examine the Computer-to-Computer Link facility",
         "its architecture, approval requirements, order routing, user and terminal "
         "identification, pre-trade risk controls and audit trail, and how it differs from the "
         "exchange-provided terminal."),
        ("To study risk management on a dealing desk",
         "margin computation in both segments, exposure and client-wise limits, margin-shortfall "
         "identification, square-off discipline and post-trade reconciliation."),
        ("To conduct a fundamental analysis of Pine Labs Limited",
         "covering its business model, industry, four-year financial record, ratio analysis, "
         "shareholding, market data, peer position, strengths and risks, without arriving at any "
         "buy or sell recommendation."),
    ]:
        R.bullet(lead, txt)

    R.section("1.3 Scope of the study")
    R.body(
        "The study covers the cash and equity-derivatives segments of the Indian secondary market "
        "as they operated during the internship period and as the rules stood on " + TODAY + ". "
        "Geographically the scope is the two principal domestic exchanges, their clearing "
        "corporations and the two depositories. Functionally it is the dealing desk and the "
        "processes immediately adjacent to it, namely risk management, back-office settlement "
        "and compliance. The currency, commodity and debt segments are outside the scope, as "
        "are portfolio management and investment advisory.")
    R.body(
        "The company analysis is confined to Pine Labs Limited and a small set of comparable "
        "listed entities. The financial period covered is FY2022-23 to FY2025-26 together with "
        "the quarter ended 30 June 2026, the latest reported period at the time of writing. "
        "Because the company listed only in November 2025, the earlier years are drawn from "
        "restated consolidated information in the offer document rather than from statutory "
        "accounts of a listed entity, and the two bases are not strictly identical. All "
        "firm-specific operational detail has been generalised in accordance with the "
        "confidentiality undertaking described in Section 4.6, and every trading example in "
        "Chapter 5 is illustrative.")

    R.section("1.4 Overview of the Indian equity and derivatives market")
    R.body(
        "The Indian secondary market is an order-driven, fully electronic market operating "
        "through two principal exchanges. Price discovery takes place by anonymous matching on "
        "**price-time priority**, with no designated market maker in the cash segment for most "
        "securities. Benchmark levels at the end of September 2026 illustrate the current state "
        "of the market: the Nifty 50 closed at " + inr(C.NIFTY_CLOSE) + " and the S&P BSE Sensex "
        "at " + inr(C.SENSEX_CLOSE) + " on 30 September 2026, after monthly declines of " +
        pct(abs(C.NIFTY_SEP_CHG)) + " per cent and " + pct(abs(C.SENSEX_SEP_CHG)) + " per cent "
        "respectively. A dealing desk experiences such a month as a period of elevated call "
        "volume, more frequent margin alerts and a higher incidence of square-off instructions, "
        "which is precisely why the operational discipline described in this report matters most "
        "when the market is least comfortable.")
    R.body(
        "Participation has broadened substantially. The two depositories together serviced about " +
        f"{C.DEMAT_TOTAL:.2f}" + " crore demat accounts as on 31 August 2026, while the National "
        "Stock Exchange reported about " + f"{C.NSE_UNIQUE_INVESTORS:.1f}" + " crore unique "
        "registered investors, measured by distinct permanent account numbers, as of July 2026. "
        "The gap arises because one investor may hold accounts with several depository "
        "participants, a distinction a dealer learns quickly, because limits and margins attach "
        "to the client code rather than to the person. The derivatives segment has grown faster "
        "than the cash segment and now dominates turnover, which has drawn sustained regulatory "
        "attention. Contract sizes have been raised, weekly expiries restricted, and the expiry "
        "days of the two exchanges separated, with National Stock Exchange equity derivatives "
        "expiring on **Tuesday** and BSE equity derivatives on **Thursday** from 1 September "
        "2025. A dealer handling clients active on both exchanges therefore manages two distinct "
        "expiry routines in the same week.")

    R.section("1.5 Challenges faced during the internship")
    R.body("The following difficulties were encountered and are recorded because they shaped both "
           "the learning and the limitations of this study.")
    for lead, txt in [
        ("Speed of the live environment",
         "order entry during volatile periods allows no time for deliberation, and the "
         "transition from classroom pace to market pace required several weeks of supervised "
         "practice before accuracy became reliable."),
        ("Volume of regulatory detail",
         "margin rules, surveillance categories, price bands and charge rates are numerous, "
         "interlocking and periodically revised, and an error in any one of them produces a real "
         "consequence rather than a lost mark."),
        ("Confidentiality constraints",
         "client-identifying data, internal limits and the firm's own financial information could "
         "not be recorded or reproduced, which required every example in this report to be "
         "reconstructed in illustrative form."),
        ("Reconciling theory with practice",
         "textbook descriptions of margining and settlement proved to be simplified, and the "
         "actual sequence of upfront margin collection, peak margin reporting and collateral "
         "pledging had to be learned from circulars and from colleagues."),
    ]:
        R.bullet(lead, txt)

    R.section("1.6 Structure of the report")
    R.body(
        "The report is organised into eight chapters and three annexures. **Chapter 1** sets out "
        "the need, objectives and scope of the study and records the challenges faced. "
        "**Chapter 2** describes Rosy Blue Securities Pvt. Ltd. and the intern's role as an "
        "Equity Dealer. **Chapter 3** is the theoretical framework and the longest chapter; it "
        "proceeds from the structure of the financial system through the trading mechanism, "
        "clearing and settlement, risk management, derivatives and margining to the NEAT PLUS "
        "terminal, the Computer-to-Computer Link facility, transaction costs and surveillance. "
        "**Chapter 4** states the research design, data, tools and ethical safeguards. "
        "**Chapter 5** contains the analysis, including the dealing-desk day, the life of an "
        "order, worked examples on futures, options, margin and charges, and the company "
        "analysis of Pine Labs Limited. **Chapters 6, 7 and 8** present the conclusion and "
        "suggestions, the limitations and the bibliography, and the three annexures give a "
        "glossary, a formula sheet and a list of abbreviations.")


# ============================================================ CHAPTER 2
def chapter2():
    R.chapter("CHAPTER 2", "ORGANISATIONAL BACKGROUND")

    R.section("2.1 Introduction to Rosy Blue Securities Pvt. Ltd.")
    R.body(
        "**Rosy Blue Securities Pvt. Ltd.** is the organisation at which the internship was "
        "undertaken. The firm operates as a **stock broker** in the Indian securities market, "
        "which means it is registered with the Securities and Exchange Board of India, holds "
        "membership of a recognised stock exchange and executes transactions in securities on "
        "behalf of clients for a brokerage commission. The specific corporate particulars are a "
        "matter of record with the Registrar of Companies and the exchanges, and have been left "
        "for completion from the firm's own documents rather than reconstructed: "
        "[ADD: year of incorporation], [ADD: registered office address], "
        "[ADD: approximate number of branches, dealing desks and employees], "
        "[ADD: names and designations of key management personnel], and "
        "[ADD: group affiliation, if any \u2014 note that no relationship between Rosy Blue "
        "Securities Pvt. Ltd. and any similarly named group has been assumed or stated].")
    R.body(
        "A broking firm of this type performs four functions that are separated for regulatory "
        "and control reasons. The **dealing function** receives and executes client orders; the "
        "**risk management function** sets and monitors client-wise exposure against collateral; "
        "the **back-office function** computes obligations, issues contract notes and reconciles "
        "against exchange and depository records; and the **compliance function** maintains "
        "client registration documents, unique client code mappings and regulatory filings. The "
        "intern was attached to the first but interacted with the other three daily, because an "
        "instruction that fails a risk check or produces an unreconciled trade becomes an "
        "immediate problem for the desk.")

    R.section("2.2 Registrations, memberships and segments")
    R.body("The firm's regulatory identity is summarised in Table 2.1. Every entry in this table "
           "is a matter of public record and should be reproduced exactly from the firm's own "
           "registration certificates and the exchange member directory; none of the "
           "identifiers has been inferred.")
    R.caption("Table 2.1: Registrations, exchange memberships and segments of the firm")
    R.table(
        ["Particular", "Entry to be confirmed from the firm's records"],
        [["Name of the entity", "Rosy Blue Securities Pvt. Ltd."],
         ["Constitution", "Private limited company incorporated under the Companies Act"],
         ["Year of incorporation", "[ADD: year of incorporation]"],
         ["Registered office", "[ADD: registered office address]"],
         ["SEBI stock-broker registration no.", "[ADD: SEBI registration number, INZ format]"],
         ["Other SEBI registrations", "[ADD: research analyst / investment adviser / DP, if held]"],
         ["Exchange membership \u2014 NSE", "[ADD: member code and segments, if a member]"],
         ["Exchange membership \u2014 BSE", "[ADD: member code and segments, if a member]"],
         ["Other exchange memberships", "[ADD: MCX / NCDEX / MSE, if applicable]"],
         ["Segments active", "[ADD: cash / equity derivatives / currency / commodity]"],
         ["Clearing arrangement", "[ADD: self-clearing or through a clearing member \u2014 name]"],
         ["Depository participation", "[ADD: CDSL / NSDL DP registration, or 'not a DP']"],
         ["Investor-grievance channel", "[ADD: e-mail and SCORES registration details]"]],
        widths=[3300, 5700], size=10.5, align=["left", "left"])
    R.source("Source: To be completed from the firm's registration certificates and the exchange "
             "member directory. No identifier in this table has been assumed by the author.")
    R.interpretation(
        "The registration and membership profile determines the boundary of everything a dealer "
        "may lawfully do. A member registered only in the cash segment cannot accept a derivatives "
        "order however insistent the client; a firm that is not a depository participant must "
        "route demat instructions through a separate entity, which lengthens the settlement "
        "chain; and a firm clearing through another member rather than self-clearing faces an "
        "additional layer in the margin and collateral process described in Section 3.9. The "
        "table is therefore not an administrative formality but the first thing a new dealer "
        "should read, because it defines the permissible universe of instructions.")

    R.section("2.3 Services offered")
    R.body(
        "The service range of an Indian retail-facing broking firm typically comprises the "
        "following, and the firm's own brochure should be used to confirm which of these it "
        "actually provides: [ADD: confirm the services offered by the firm].")
    for lead, txt in [
        ("Secondary-market broking",
         "execution of buy and sell orders in the cash segment on behalf of registered clients, "
         "against a brokerage charged on turnover or on a per-order basis."),
        ("Derivatives broking",
         "execution in equity futures and options, subject to the client having satisfied the "
         "additional documentation and margin requirements attaching to that segment."),
        ("Depository services",
         "holding of client securities in dematerialised form, where the firm is a depository "
         "participant, together with pledge, re-pledge and transfer instructions."),
        ("Research and advisory",
         "dissemination of research output, which may be undertaken only to the extent permitted "
         "by the firm's separate registration as a research analyst or investment adviser."),
    ]:
        R.bullet(lead, txt)

    R.section("2.4 Organisation structure and the dealing desk")
    R.body("The functional organisation of a broking firm and the position of the dealing desk "
           "within it are shown in Figure 2.1, and the responsibilities of each department are "
           "set out in Table 2.2.")
    R.caption("Figure 2.1: Organisation structure and the position of the dealing desk")
    R.image(G.fig_2_1())
    R.source("Source: Compiled by the author; the exact structure of the firm is to be confirmed.")
    R.caption("Table 2.2: Departments of a broking firm and their principal functions")
    R.table(
        ["Department", "Principal functions", "Interaction with the dealing desk"],
        [["Dealing desk", "Receives client instructions, enters and modifies orders, confirms "
                          "executions, communicates fills to clients",
          "The function itself"],
         ["Risk management (RMS)", "Sets client-wise limits, monitors margin utilisation, raises "
                                   "shortfall alerts, authorises square-off",
          "Releases or blocks limits in real time during market hours"],
         ["Back office and settlement", "Computes obligations, generates contract notes, "
                                        "maintains client ledgers, handles pay-in and pay-out",
          "Supplies opening ledger balances and resolves trade mismatches"],
         ["Compliance", "Client registration, unique client code mapping, surveillance responses, "
                        "regulatory filings and reporting",
          "Confirms that a client code is active and permitted in a segment"],
         ["Information technology", "Terminal provisioning, connectivity, user and terminal "
                                    "identifiers, business-continuity arrangements",
          "Restores terminal access and connectivity when a session fails"],
         ["Accounts and treasury", "Client fund transfers, bank reconciliation, settlement "
                                   "funding and collateral placement",
          "Confirms receipt of client funds that release additional limit"]],
        widths=[1900, 4000, 3100], size=9.5, align=["left", "left", "left"])
    R.source(SRC_COMP)
    R.interpretation(
        "The separation of the dealing function from the risk-management function is the single "
        "most important control in the structure, because it prevents the person with an "
        "incentive to execute from also being the person who decides whether execution is "
        "affordable. A dealer therefore cannot increase a client's limit, and an instruction "
        "that exceeds it must be refused at the desk and escalated. During the internship this "
        "separation was the most frequent source of friction with clients and simultaneously "
        "the most frequent source of protection for the firm, which is consistent with the "
        "margin-shortfall scenario computed in Section 5.8.")

    R.section("2.5 Mission, vision and values")
    R.body("[ADD: reproduce the firm's published mission, vision and values statement here. If "
           "the firm does not publish one, delete this section entirely and renumber Section 2.6 "
           "as Section 2.5 together with the corresponding Contents entry.]")

    R.section("2.6 The intern's role as an equity dealer")
    R.body(
        "The intern was attached to the equity dealing desk in the capacity of an **Equity "
        "Dealer** for the period [ADD: internship start date] to [ADD: internship end date], "
        "working [ADD: working hours] and reporting to [ADD: designation of the reporting "
        "senior, role only, no names]. The segments handled were [ADD: cash / equity derivatives "
        "/ both] and the client categories dealt with were [ADD: retail / high net-worth / "
        "sub-broker / institutional \u2014 categories only]. All order entry was carried out "
        "under the supervision of a registered dealer, and no instruction was executed without a "
        "verified client instruction and a confirmed limit.")
    R.body(
        "The responsibilities fell into three time bands, corresponding to the periods before, "
        "during and after the continuous trading session, and are set out in Table 2.3.")
    R.caption("Table 2.3: Indicative working day of an equity dealer")
    R.table(
        ["Time", "Activity", "Purpose"],
        [["08:45 \u2013 09:00", "Review of overnight news, corporate actions, the exchange ban "
                                "list and pending client instructions",
          "To identify scrips that require special handling before the session opens"],
         ["09:00 \u2013 09:07", "Verification of client-wise limits and collateral uploaded by "
                                "the risk-management system",
          "To confirm the tradable limit available to each client for the day"],
         ["09:07 \u2013 09:15", "Observation of the pre-open session and the indicative "
                                "equilibrium price",
          "To assess the likely opening level"],
         ["09:15 \u2013 12:30", "Continuous trading: order entry, modification and cancellation; "
                                "confirmation of fills to clients",
          "Execution of client instructions within the limits available"],
         ["12:30 \u2013 14:30", "Mid-session monitoring, response to risk alerts, margin top-up "
                                "follow-up",
          "To prevent a margin shortfall from becoming a forced square-off"],
         ["14:30 \u2013 15:30", "Intraday square-off discipline, closing instructions, "
                                "end-of-session order review",
          "To close positions that cannot be carried forward under the client's limit"],
         ["15:30 \u2013 17:30", "Verification of the day's executions against the trade book, "
                                "post-trade reconciliation with the back office and "
                                "contract-note verification",
          "To ensure the client's records and the firm's records agree"]],
        widths=[1500, 4000, 3500], size=9.5, align=["center", "left", "left"])
    R.source("Source: Author's observation during the internship; timings are indicative of a "
             "normal trading day and exclude special sessions.")
    R.interpretation(
        "The shape of the day shows that only about six of the nine working hours fall inside "
        "the continuous trading session, and that much of a dealer's effort is expended before "
        "the market opens and after it closes. The pre-market block determines how much can "
        "safely be executed; the post-market block determines whether what was executed has "
        "been recorded correctly. Order entry is thus the visible part of a process that is "
        "predominantly preparation and verification, consistent with the reconciliation "
        "emphasis described in Section 3.17.")
    R.body("The principal tools used during the internship were the following.")
    for lead, txt in [
        ("The exchange-provided trading terminal",
         "used for market watch, order entry, order-book and trade-book review, as described in "
         "Section 3.15."),
        ("The member's Computer-to-Computer Link terminal",
         "used for order routing through the firm's own server with pre-trade risk controls "
         "applied, as described in Section 3.16. "
         "[ADD: name of the CTCL software used by the firm]."),
        ("The risk-management system console",
         "used to read client-wise limit utilisation, available margin and shortfall alerts."),
        ("The telephone recording system",
         "used for the receipt of client instructions, retained as evidence of the instruction in "
         "accordance with the firm's compliance policy."),
    ]:
        R.bullet(lead, txt)


# ============================================================ ASSEMBLY
import part2, part3, part4

def main():
    cover()
    certificate()
    experience_cert()
    declaration()
    acknowledgement()
    executive_summary()
    contents()
    chapter1()
    chapter2()
    part2.chapter3(R)
    part2.chapter4(R)
    part3.chapter5(R)
    part4.chapter6(R)
    part4.chapter7(R)
    part4.chapter8(R)
    part4.annexures(R)
    R.save(OUT)
    print("saved:", OUT)


if __name__ == "__main__":
    main()
