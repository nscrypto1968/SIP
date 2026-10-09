#!/usr/bin/env python3
"""Consistency checks on the built document."""
import re, sys, collections
from docx import Document
from docx.shared import Twips

PATH = "/home/user/SIP/Rosy_Blue_Securities_Black_Book.docx"
doc = Document(PATH)

words = 0
caps_t, caps_f = [], []
refs_t, refs_f = set(), set()
placeholders = []
tbv = 0
srcs = 0
for p in doc.paragraphs:
    t = p.text
    words += len(t.split())
    m = re.match(r"^Table ([A-C]|\d+)\.(\d+):", t)
    if m:
        caps_t.append(f"{m.group(1)}.{m.group(2)}")
    m = re.match(r"^Figure (\d+)\.(\d+):", t)
    if m:
        caps_f.append(f"{m.group(1)}.{m.group(2)}")
    for r in re.findall(r"Table ([A-C]|\d+)\.(\d+)", t):
        refs_t.add(f"{r[0]}.{r[1]}")
    for r in re.findall(r"Figure (\d+)\.(\d+)", t):
        refs_f.add(f"{r[0]}.{r[1]}")
    placeholders += re.findall(r"\[ADD:[^\]]*\]", t)
    if "to be verified" in t:
        tbv += 1
    if t.startswith("Source:"):
        srcs += 1

tw = 0
for tb in doc.tables:
    for row in tb.rows:
        for c in row.cells:
            tw += len(c.text.split())
            placeholders += re.findall(r"\[ADD:[^\]]*\]", c.text)
            if "to be verified" in c.text:
                pass

print("paragraph words :", words)
print("table words     :", tw)
print("TOTAL words     :", words + tw)
print("paragraphs      :", len(doc.paragraphs))
print("tables          :", len(doc.tables))
print("inline images   :", len(doc.inline_shapes))
print("source lines    :", srcs)
print("'to be verified':", tbv)

def check(name, caps):
    dup = [k for k, v in collections.Counter(caps).items() if v > 1]
    print(f"\n{name} captions ({len(caps)}): {caps}")
    if dup:
        print("  !! DUPLICATES:", dup)
    bych = collections.defaultdict(list)
    for c in caps:
        ch, n = c.split(".")
        bych[ch].append(int(n))
    for ch, ns in bych.items():
        exp = list(range(1, len(ns) + 1))
        if sorted(ns) != exp:
            print(f"  !! chapter {ch} sequence broken: {sorted(ns)} expected {exp}")
        else:
            print(f"  chapter {ch}: 1..{len(ns)} OK")

check("TABLE", caps_t)
check("FIGURE", caps_f)

missing_t = sorted(set(caps_t) - refs_t)
missing_f = sorted(set(caps_f) - refs_f)
print("\ncaptioned but never referenced in text  - tables:", missing_t or "none")
print("captioned but never referenced in text  - figures:", missing_f or "none")
orph_t = sorted(refs_t - set(caps_t))
orph_f = sorted(refs_f - set(caps_f))
print("referenced but no caption exists - tables:", orph_t or "none")
print("referenced but no caption exists - figures:", orph_f or "none")

# table width integrity
bad = 0
for i, tb in enumerate(doc.tables):
    gw = [int(g.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w'))
          for g in tb._tbl.findall(
              '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblGrid/'
              '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}gridCol')]
    if gw and sum(gw) not in (9000, 2400):
        print(f"  !! table {i} grid sums to {sum(gw)}")
        bad += 1
print("\ntables with non-conforming width:", bad)

print("\nyellow placeholders:", len(placeholders))
seen = []
for p in placeholders:
    if p not in seen:
        seen.append(p)
for p in seen:
    print("   ", p)
print("unique placeholders:", len(seen))
